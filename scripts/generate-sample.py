"""Generate public/sample-chart.xlsx: a bar chart whose series carry EXPLICIT,
distinct solid-fill colors (red / green / blue).

Explicit per-series colors are used on purpose: they must survive any faithful
round-trip, so if they go missing after the viewer re-serializes the workbook
there is no "it fell back to a theme default" ambiguity — it is a clear loss.
The legend swatches derive from these series colors, so legend color loss is the
visible symptom.

Run: .venv/bin/python scripts/generate-sample.py
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference, Series
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.fill import PatternFillProperties, ColorChoice

OUT = Path(__file__).resolve().parent.parent / "public" / "sample-chart.xlsx"

SERIES = [
    ("Alpha", [5, 8, 6, 9], "FF0000"),   # red
    ("Beta", [3, 4, 7, 2], "00B050"),    # green
    ("Gamma", [6, 2, 5, 8], "0070C0"),   # blue
]
CATEGORIES = ["Q1", "Q2", "Q3", "Q4"]


def main() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"

    ws.append(["Quarter", *[name for name, _, _ in SERIES]])
    for row_idx, cat in enumerate(CATEGORIES):
        ws.append([cat, *[values[row_idx] for _, values, _ in SERIES]])

    chart = BarChart()
    chart.type = "col"
    chart.title = "Quarterly results (explicit series colors)"
    chart.style = 10

    cats = Reference(ws, min_col=1, min_row=2, max_row=1 + len(CATEGORIES))
    for col_idx, (name, _values, hex_color) in enumerate(SERIES, start=2):
        data = Reference(ws, min_col=col_idx, min_row=1, max_row=1 + len(CATEGORIES))
        series = Series(data, title_from_data=True)
        series.graphicalProperties = GraphicalProperties(
            solidFill=ColorChoice(srgbClr=hex_color)
        )
        chart.series.append(series)
    chart.set_categories(cats)

    ws.add_chart(chart, "F2")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"wrote {OUT} ({OUT.stat().st_size} bytes)")
    print("series colors: " + ", ".join(f"{n}={c}" for n, _, c in SERIES))


if __name__ == "__main__":
    main()
